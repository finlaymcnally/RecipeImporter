from __future__ import annotations

from pathlib import Path

import docx

from cookimport.config.run_settings import RunSettings
from cookimport.plugins.text import TextImporter, _extract_sections_from_blob
from tests.paths import FIXTURES_DIR as TESTS_FIXTURES_DIR


FIXTURES_DIR = TESTS_FIXTURES_DIR


def _source_text(result) -> str:
    return "\n".join(block.text for block in result.source_blocks)


def test_detect_text_file():
    importer = TextImporter()
    assert importer.detect(Path("recipe.txt")) > 0.8
    assert importer.detect(Path("recipe.md")) > 0.8
    assert importer.detect(Path("image.png")) == 0.0


def test_inspect_simple_text():
    importer = TextImporter()
    inspection = importer.inspect(FIXTURES_DIR / "simple_text.txt")
    assert len(inspection.sheets) == 1
    assert inspection.sheets[0].layout == "single-recipe"


def test_inspect_multi_recipe():
    importer = TextImporter()
    inspection = importer.inspect(FIXTURES_DIR / "multi_recipe.md")
    assert len(inspection.sheets) == 1
    assert any("2 recipe candidate(s)" in w for w in inspection.sheets[0].warnings)


def test_convert_simple_text():
    importer = TextImporter()
    result = importer.convert(FIXTURES_DIR / "simple_text.txt", None)

    assert result.recipes == []
    assert len(result.source_blocks) >= 1
    assert len(result.source_support) == 1
    assert "Simple Pasta" in _source_text(result)
    assert "1 lb Pasta" in _source_text(result)
    assert "Boil water." in _source_text(result)


def test_convert_multi_recipe():
    importer = TextImporter()
    result = importer.convert(FIXTURES_DIR / "multi_recipe.md", None)

    assert result.recipes == []
    assert len(result.source_support) == 2
    assert "Recipe One" in _source_text(result)
    assert "Item A" in _source_text(result)
    assert "Recipe Two" in _source_text(result)
    assert "Item C" in _source_text(result)


def test_convert_multi_recipe_rules_v1_backend() -> None:
    importer = TextImporter()
    result = importer.convert(
        FIXTURES_DIR / "multi_recipe.md",
        None,
        run_settings=RunSettings(
            multi_recipe_splitter="rules_v1",
        ),
    )

    assert result.recipes == []
    assert len(result.source_support) == 2


def test_convert_multi_recipe_splitter_off_keeps_single_candidate() -> None:
    importer = TextImporter()
    result = importer.convert(
        FIXTURES_DIR / "multi_recipe.md",
        None,
        run_settings=RunSettings(multi_recipe_splitter="off"),
    )

    assert result.recipes == []
    assert len(result.source_support) == 1


def test_convert_rules_v1_guardrail_keeps_for_the_component_with_recipe(
    tmp_path: Path,
) -> None:
    source = tmp_path / "for_the_guardrail.txt"
    source.write_text(
        "\n".join(
            [
                "Cake",
                "Ingredients",
                "1 cup flour",
                "Instructions",
                "Bake cake.",
                "For the frosting",
                "1 tbsp butter",
                "Stir until creamy.",
                "Cookie",
                "Ingredients",
                "2 eggs",
                "Instructions",
                "Whisk eggs.",
            ]
        ),
        encoding="utf-8",
    )
    importer = TextImporter()
    result = importer.convert(
        source,
        None,
        run_settings=RunSettings(multi_recipe_splitter="rules_v1"),
    )
    assert result.recipes == []
    assert len(result.source_support) == 2
    assert "Cake" in _source_text(result)
    assert "Cookie" in _source_text(result)


def test_convert_serves_split_text():
    importer = TextImporter()
    result = importer.convert(FIXTURES_DIR / "serves_multi.txt", None)

    assert result.recipes == []
    assert len(result.source_support) == 2
    assert "Simple Salad" in _source_text(result)
    assert "Serves 2" in _source_text(result)
    assert "Quick Eggs" in _source_text(result)
    assert "Serves 4" in _source_text(result)


def test_convert_docx_table(tmp_path: Path):
    doc = docx.Document()
    table = doc.add_table(rows=3, cols=3)
    headers = ["Recipe", "Ingredients", "Instructions"]
    for idx, value in enumerate(headers):
        table.cell(0, idx).text = value
    table.cell(1, 0).text = "Table Recipe One"
    table.cell(1, 1).text = "1 cup rice\n2 cups water"
    table.cell(1, 2).text = "Rinse rice.\nCook until tender."
    table.cell(2, 0).text = "Table Recipe Two"
    table.cell(2, 1).text = "- 2 eggs\n- 1 tbsp butter"
    table.cell(2, 2).text = "Beat eggs.\nCook in butter."
    path = tmp_path / "table_recipes.docx"
    doc.save(path)

    importer = TextImporter()
    result = importer.convert(path, None)

    assert result.recipes == []
    assert len(result.source_blocks) == 2
    assert len(result.source_support) == 2
    assert result.source_support[0].kind == "docx_table_row"
    assert "Table Recipe One" in _source_text(result)
    assert "1 cup rice" in _source_text(result)
    assert "Rinse rice." in _source_text(result)


def test_text_importer_preserves_split_candidates_as_non_authoritative_support(
    tmp_path: Path,
) -> None:
    source = tmp_path / "mixed.txt"
    source.write_text(
        "\n".join(
            [
                "Skillet Beans",
                "Ingredients",
                "1 can beans",
                "1 tbsp olive oil",
                "Instructions",
                "Heat oil in a pan.",
                "Add beans and warm through.",
                "",
                "==== RECIPE ====",
                "",
                "Kitchen Notes",
                "This section talks about pantry strategy and chapter context only.",
            ]
        ),
        encoding="utf-8",
    )

    result = TextImporter().convert(source, None)

    assert result.recipes == []
    assert len(result.source_support) == 2
    assert all(item.kind == "candidate_recipe_region" for item in result.source_support)
    assert "Skillet Beans" in _source_text(result)
    assert "Kitchen Notes" in _source_text(result)
    assert not any(
        artifact.location_id == "recipe_scoring_debug" for artifact in result.raw_artifacts
    )


def test_text_blob_section_extraction_keeps_for_component_headers() -> None:
    sections = _extract_sections_from_blob(
        "\n".join(
            [
                "Ingredients",
                "For the sauce:",
                "1 tbsp butter",
                "Instructions",
                "For the sauce:",
                "Melt the butter.",
            ]
        )
    )
    assert sections["ingredients"] == ["For the sauce", "1 tbsp butter"]
    assert sections["instructions"] == ["For the sauce", "Melt the butter."]
